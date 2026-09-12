#!/usr/bin/env python3
"""
Build FPC training document inventory from docs/training/.

Outputs:
  - exports/fpc-knowledge/document_inventory.json
  - docs/FPC_TRAINING_DOCUMENT_INDEX.md

Scans corpus roots:
  FPC Documents 1/, FPC-Docs 2/, FPC Docs 3/, P&A Documents/
Notes sibling *.zip archives. Extracts ASC/table names from .docx via
python-docx when available (else zip/XML), validated against an optional
RUN ASC catalog. Priority CRITICAL/HIGH docs carry curated reviewed metadata.
"""
from __future__ import annotations

import argparse
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
TRAINING_ROOT = REPO_ROOT / "docs" / "training"
DEFAULT_JSON = REPO_ROOT / "exports" / "fpc-knowledge" / "document_inventory.json"
DEFAULT_MD = REPO_ROOT / "docs" / "FPC_TRAINING_DOCUMENT_INDEX.md"
DEFAULT_ASC_DIR = REPO_ROOT / "workspace" / "active" / "RUN" / "FORTNA"

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DOC_EXTS = {".docx", ".doc", ".pdf", ".xlsx", ".xlsm", ".pptx", ".txt"}
CORPUS_DIRS = (
    "FPC Documents 1",
    "FPC-Docs 2",
    "FPC Docs 3",
    "P&A Documents",
)

RELEVANCE_ORDER = {
    "CRITICAL": 0,
    "HIGH": 1,
    "MEDIUM": 2,
    "LOW": 3,
    "OUT_OF_SCOPE": 4,
}

# Stopwords / weak ASC stems to drop unless explicitly curated.
WEAK_ASC = {
    "action",
    "area",
    "belt",
    "blank",
    "colors",
    "config",
    "contact",
    "groups",
    "keeping",
    "logic",
    "messages",
    "operator",
    "opt",
    "optimizer",
    "parameters",
    "password",
    "product",
    "state",
    "updates",
    "timemenu",
    "horns",
    "batches",  # often incidental; keep when curated
}

# Reject non-ASC tokens that appear as false-positive "X Table" matches.
REJECT_TABLE_NAMES = WEAK_ASC | {
    "the",
    "this",
    "that",
    "from",
    "with",
    "each",
    "following",
    "file",
    "files",
    "name",
    "type",
    "types",
    "map",
    "list",
    "full",
    "add",
    "use",
    "whole",
    "down",
    "lane",
    "merge",
    "scan",
    "sort",
    "host",
    "detail",
    "level",
    "message",
    "tracking",
    "example",
    "status",
    "address",
    "contents",
    "created",
    "details",
    "support",
    "found",
    "always",
    "indicates",
    "uses",
    "same",
    "are",
    "and",
    "for",
    "into",
    "using",
    "own",
    "first",
    "second",
    "timer",
    "refresh",
    "popup",
    "dat8",
    "crr",
    "dcm",
    "mac",
    "run",
    "fpc",
    "bit",
    "ctrl",
}

TABLE_PAT = re.compile(r"\b([A-Za-z][A-Za-z0-9_]{1,40})\s+Table\b")
ASC_PAT = re.compile(r"\b([A-Za-z][A-Za-z0-9_]{1,40})\.(?:asc|ASC)\b")
WORD_PAT = re.compile(r"\b([A-Za-z][A-Za-z0-9_]{2,40})\b")

# ---------------------------------------------------------------------------
# Curated reviewed metadata for priority CRITICAL/HIGH docs (deep skim).
# tables_mentioned use exact ASC-like names from the docs (not invented).
# ---------------------------------------------------------------------------
REVIEWED: dict[str, dict] = {
    "FPC-Merge-Modules": {
        "relevance": "CRITICAL",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "merge",
        "tables_mentioned": [
            "SimpleMerge",
            "MergeBoss",
            "MergeInputs",
            "MergeRoute",
            "MergeNotBusy",
            "MergeRunOutputs",
            "MergeStopOutputs",
            "SawMerge",
            "SawLane",
            "ZipperMerge",
            "ZipperLane",
            "ServoGapper",
            "ZprSimConfig",
            "ZipperErrLog",
            "ZprIOAction",
            "ZprStatusMsgs",
            "ZprTickGapper",
            "ZprGapperZone",
            "Fullline",
            "Inpoints",
            "RateCount",
            "Batches",
        ],
        "modules_mentioned": [
            "Simple Merge",
            "Multiple Input Lane Merge (New Merge)",
            "Sawtooth Merge",
            "High-Speed Sawtooth Merge",
            "Zipper Merge",
            "Zipper Merge Tick Gapper",
            "Servo Gapper",
            "Inch-and-Store",
        ],
        "relationships_mentioned": [
            "Sawtooth Merge uses SawMerge + SawLane with merge shifter reservations",
            "New Merge links MergeBoss to MergeInputs / MergeRoute / MergeNotBusy / run-stop outputs",
            "Zipper Merge uses ZipperMerge + ZipperLane and optional ServoGapper / ZprTickGapper",
            "High-Speed Sawtooth Merge covered in dedicated HSSaw* document",
        ],
        "reviewed": True,
        "notes": "Primary merge corpus: Simple/New/Sawtooth/Zipper. Critical for PLC merge generation and ASC import.",
    },
    "FPC-HighSpeedSawtoothMerge": {
        "relevance": "CRITICAL",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "merge/sawtooth",
        "tables_mentioned": [
            "HSSawMerge",
            "HSSawLane",
            "HSSawParm",
            "HSSawSim",
            "SawLane",
            "SawMerge",
            "Fullline",
            "Inpoints",
            "SortBuff",
            "Batches",
            "Conveyor",
        ],
        "modules_mentioned": [
            "High-Speed Sawtooth Merge",
            "Merge Shifter",
            "Inch-and-Store",
        ],
        "relationships_mentioned": [
            "HSSawMerge owns merge belt; HSSawLane defines lanes; HSSawParm holds per-lane parameter sets",
            "HSSawSim supports simulation/test; lane PE / full timer / priority IO select active HSSawParm rows",
            "Uses standard shifter linked to merge belt with overlapping induction points",
        ],
        "reviewed": True,
        "notes": "Canonical HS sawtooth ASC model (HSSawMerge/Lane/Parm/Sim). Aligns with Site Forge Sawtooth Merge generation.",
    },
    "FPC-Sorter-Control-Module": {
        "relevance": "CRITICAL",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "sorter",
        "tables_mentioned": [
            "Sorters",
            "SrtAppControl",
            "SrtScanBoss",
            "SrtZoneLane",
            "SrtBadGapCnfg",
            "SrtLaneNotAvail",
            "SrtRndRobin",
            "SrtTrack1",
            "SrtDevice1",
            "ScnScanDevice",
            "ScnScanZone",
            "ScnDeviceErrIO",
            "ScnZoneErrIO",
            "MsgMap",
            "MsgCR1",
            "MsgDC1",
            "MsgGP1",
            "Machine",
            "Outpoints",
            "SortData",
            "Batches",
            "Errors",
        ],
        "modules_mentioned": [
            "Sorter Control",
            "Scanner Control",
            "Host messaging (CRR/DCM/GP)",
        ],
        "relationships_mentioned": [
            "SrtScanBoss links scan zones to sorters; ScnScanDevice feeds tracking / ProcFlag",
            "SrtAppControl selects message tables; MsgMap + Machine define host connections",
            "SrtZoneLane maps sorter zones to divert lanes",
        ],
        "reviewed": True,
        "notes": "Core sorter control/reference. Large module covering scan, divert, host messaging, tracking tables.",
    },
    "FPC-Shifter-And-Sorter-Configuration": {
        "relevance": "CRITICAL",
        "document_type": "howto",
        "generation_relevance": "PLC_GEN",
        "subsystem": "sorter/shifter",
        "tables_mentioned": [
            "Sorters",
            "SortBuff",
            "SortData",
            "SrtBadGapCnfg",
            "SrtZoneLane",
            "ScnScanDevice",
            "Conveyor",
            "Encoders",
            "Inpoints",
            "Outpoints",
            "RateCount",
            "Jamzones",
            "Errors",
            "Machine",
        ],
        "modules_mentioned": [
            "Shifter",
            "Sorter",
            "Scanner",
        ],
        "relationships_mentioned": [
            "Shifter configuration ties Conveyor / Encoders / Inpoints to sorter tracking",
            "SortBuff / SortData hold tracking state used by divert logic",
            "References HSSawMerge context when merges feed sorters",
        ],
        "reviewed": True,
        "notes": "How-to for shifter + sorter configuration; complements Sorter Control Module.",
    },
    "FPC-SorterConfigurationChecklist": {
        "relevance": "HIGH",
        "document_type": "checklist",
        "generation_relevance": "PLC_GEN",
        "subsystem": "sorter",
        "tables_mentioned": [
            "Batches",
            "SrtAppControl",
            "ScnScanZone",
            "ScnScanDevice",
            "SrtScanBoss",
            "Machine",
            "MsgMap",
            "SrtZoneLane",
            "MsgTST",
            "AsciiMnu",
            "Outpoints",
        ],
        "modules_mentioned": [
            "Sorter Control",
            "Scanner Control",
            "CommCore messaging",
        ],
        "relationships_mentioned": [
            "Scan event flow: Machine/MsgMap -> ScnScanDevice -> SrtScanBoss -> SrtAppControl message table -> MsgMap TX",
            "Checklist companion to FPC-Sorter-Control-Module",
        ],
        "reviewed": True,
        "notes": "Operational checklist for sorter bring-up; high value for generation completeness gates.",
    },
    "Unit Sorter Control Software V1 rev25 DDD": {
        "relevance": "HIGH",
        "document_type": "ddd",
        "generation_relevance": "PLC_GEN",
        "subsystem": "sorter/unit_sorter",
        "tables_mentioned": [
            "SortData",
            "Sorters",
            "Groups",
            "Machine",
            "Outpoints",
            "Conveyor",
            "UnitScanBoss",
            "UnitSrtAppCtrl",
            "UnitUpdPnts",
        ],
        "modules_mentioned": [
            "Unit Sorter Control",
            "Host messaging (UII/UDI/UCM)",
            "Chute Control",
            "Ethernet IP Scan Interface",
            "Scanner Statistics",
        ],
        "relationships_mentioned": [
            "DDD for unit sorter platform; unit-specific tables (UnitScanBoss/UnitSrtAppCtrl/UnitUpdPnts) alongside SortData/Sorters",
            "Covers induction, chute assignment, heartbeat, exception, and speed-check behaviors",
        ],
        "reviewed": True,
        "notes": "Legacy .doc DDD (rev25). Skimmed via binary strings. Prefer rev25 over rev24. Some unit-sorter table names are DDD-specific.",
    },
    "FPC-StartStopZones": {
        "relevance": "CRITICAL",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "zones/safety",
        "tables_mentioned": [
            "StartStopZones",
            "Jamzones",
            "Jamcheck",
            "CombinedJamZones",
            "CombinedEnableBits",
            "Mtrchain",
            "Conveyor",
            "Errors",
            "Machine",
        ],
        "modules_mentioned": [
            "Start/Stop Zones",
            "Combined Jam Zones",
            "Motor Startup Chains",
        ],
        "relationships_mentioned": [
            "StartStopZones group jam zones that start/stop together",
            "Jamzones Enable Bit may be built via CombinedJamZones + CombinedEnableBits",
            "Jamzones Latch Bit links to Mtrchain Motor_Aux to start motor chains; Stop Zone can halt chain",
            "Jamcheck sensors link into Jamzones via Zones column",
        ],
        "reviewed": True,
        "notes": "Defines zone/jam/enable model required for area safety and motor chain PLC generation.",
    },
    "FPC-Fulls-Jams-Fulljams": {
        "relevance": "CRITICAL",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "zones/full_jam",
        "tables_mentioned": [
            "Fullline",
            "Jamcheck",
            "Fulljam",
            "Jamzones",
            "Conveyor",
            "Errors",
            "Batches",
        ],
        "modules_mentioned": [
            "Full Detection",
            "Jam Detection",
            "Full-Jams",
        ],
        "relationships_mentioned": [
            "Fulls use Fullline (auto-restart); Jams use Jamcheck (operator clear); Full-jams use Fulljam hybrid",
            "All rely on photoeyes + timers; display/error presentation configured per table",
            "Related to Jamzones enable/stop behavior documented in StartStopZones",
        ],
        "reviewed": True,
        "notes": "Combined fulls/jams/fulljams reference. Tables Fullline, Jamcheck, Fulljam are primary ASC names.",
    },
    "FPC-Motor-Startup-Chains": {
        "relevance": "CRITICAL",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "motors",
        "tables_mentioned": [
            "Mtrchain",
            "Conveyor",
            "Convclr",
            "Jamzones",
            "Fulljam",
            "Fullline",
            "Errors",
        ],
        "modules_mentioned": [
            "Motor Startup Chains",
            "Start/Stop Zones",
        ],
        "relationships_mentioned": [
            "Mtrchain sequences motor starts; Motor_Aux latch typically driven by Jamzones Latch Bit",
            "Stop Zone / overload paths interact with jam and fulljam conditions",
            "Conveyor / Convclr participate in clear and run presentation",
        ],
        "reviewed": True,
        "notes": "Motor chain ASC is Mtrchain. Critical linkage between jam zones and conveyor motor startup PLC.",
    },
    "FPC-The-ViewIO-Screen-and-FORTNADT-Table": {
        "relevance": "HIGH",
        "document_type": "howto",
        "generation_relevance": "IO",
        "subsystem": "io/ui",
        "tables_mentioned": [
            "FORTNADT",
            "Configio",
            "Conveyor",
            "FastIO",
            "flagmnu",
        ],
        "modules_mentioned": [
            "View I/O Screen",
            "Disable Bits",
            "FastIO",
        ],
        "relationships_mentioned": [
            "View I/O displays FORTNADT bit force/invert/on-off by decimal word row",
            "Word addresses align with Configio Octal_Word; part names from Conveyor",
            "Disable/GoNow come from Conveyor_Disable_I_O / Conveyor_Overide_I_O",
        ],
        "reviewed": True,
        "notes": "Documents FORTNADT.asc runtime bit table and View I/O debugging. Doc cites flagmnu.asc spelling.",
    },
    "FPC-Machine-MsgMap-Configuraton": {
        "relevance": "HIGH",
        "document_type": "howto",
        "generation_relevance": "SITE_MODEL",
        "subsystem": "messaging",
        "tables_mentioned": [
            "Machine",
            "MsgMap",
            "Connect",
            "Protocol",
            "AsciiMnu",
            "RecvStat",
            "SendStat",
            "MsgCR1",
            "Scanners",
            "ScnScanDevice",
        ],
        "modules_mentioned": [
            "Machine configuration",
            "MsgMap routing",
            "Host/scanner communications",
        ],
        "relationships_mentioned": [
            "Machine defines hosts/connections; MsgMap routes messages to processes/devices",
            "Scanner and CommCore connections configured via Machine + MsgMap pairs",
        ],
        "reviewed": True,
        "notes": "Filename spelling Configuraton retained. Core for multi-AC and host message topology.",
    },
    "FPC-IOCard-Interfaces": {
        "relevance": "HIGH",
        "document_type": "module_ref",
        "generation_relevance": "IO",
        "subsystem": "io",
        "tables_mentioned": [
            "IOCard",
            "IOCardResp",
            "Configio",
            "FORTNADT",
            "ABS_CARD_INFO",
            "KTX_RIO_ADAPTERS",
            "KTX_RAM",
            "Conveyor",
            "Machine",
            "Outpoints",
            "Errors",
        ],
        "modules_mentioned": [
            "IOCard Interfaces",
            "Ethernet/IP adapters",
            "RIO adapters",
        ],
        "relationships_mentioned": [
            "IOCard defines card/channel mapping into Configio / FORTNADT word space",
            "ABS_CARD_INFO and KTX_* tables support adapter inventory/status",
        ],
        "reviewed": True,
        "notes": "Hardware IO card interface reference for mapping physical IO into FPC tables.",
    },
    "FPC-FastIO-Configuration": {
        "relevance": "HIGH",
        "document_type": "howto",
        "generation_relevance": "IO",
        "subsystem": "io",
        "tables_mentioned": [
            "FastIO",
            "Configio",
            "Conveyor",
            "Machine",
        ],
        "modules_mentioned": [
            "FastIO",
            "IO Synchronization",
        ],
        "relationships_mentioned": [
            "FastIO sends Configio/Conveyor word blocks from SrcMachine to DestMachine faster than 1 Hz sync",
            "Receiving AC needs no extra FastIO rows; ownership must be on sender",
        ],
        "reviewed": True,
        "notes": "Inter-AC fast IO word replication. Not for safety-critical hardwired substitutes.",
    },
    "FPC-Scanner-Control-Configuration": {
        "relevance": "HIGH",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "scanner",
        "tables_mentioned": [
            "Scanners",
            "ScnScanDevice",
            "ScnScanZone",
            "ScnDeviceErrIO",
            "ScnZoneErrIO",
            "SrtScanBoss",
            "SrtTrack1",
            "MsgMap",
            "MsgCR1",
            "MsgDC1",
            "MsgGP1",
            "Machine",
            "Outpoints",
            "XfrDevice",
            "XfrTrack",
            "XfRouteBoss",
            "HistConfig",
            "AsciiMnu",
            "Errors",
        ],
        "modules_mentioned": [
            "Scanner Control",
            "Sorter scan boss",
            "Transfer scan routing",
        ],
        "relationships_mentioned": [
            "ScnScanZone/Device define scan points; errors via ScnDeviceErrIO / ScnZoneErrIO",
            "Sorter path uses SrtScanBoss; transfer path uses XfrDevice / XfrTrack / XfRouteBoss",
            "Host messaging via MsgMap + Machine",
        ],
        "reviewed": True,
        "notes": "Scanner configuration spanning sorter and transfer scan pipelines.",
    },
    "FPC-Gapping-Modules": {
        "relevance": "HIGH",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "gapping",
        "tables_mentioned": [
            "NewGapper",
            "NewGapperStates",
            "ZprTickGapper",
            "ZprGapperZone",
            "ServoGapper",
            "CapacityIO",
            "CapacityBoss",
            "CapacityGapCtrl",
            "RateCount",
            "Batches",
            "Conveyor",
            "Fullline",
            "ZipperLane",
        ],
        "modules_mentioned": [
            "New Gapper",
            "Tick Gapper",
            "Servo Gapper",
            "Capacity Gapper",
            "GapCfg (legacy)",
        ],
        "relationships_mentioned": [
            "NewGapper replaces GapCfg; can share discharge with merges",
            "Tick gapper uses ZprTickGapper + ZprGapperZone (encoder ticks)",
            "Capacity gapper uses CapacityIO / CapacityBoss / CapacityGapCtrl; examples with Zipper Merge",
        ],
        "reviewed": True,
        "notes": "Gapping family used by merges/sorters for slug spacing and flow control.",
    },
    "FPC-Motor-Speed-Control": {
        "relevance": "HIGH",
        "document_type": "module_ref",
        "generation_relevance": "PLC_GEN",
        "subsystem": "motors",
        "tables_mentioned": [
            "SpdControl",
            "Conveyor",
            "ZipperLane",
        ],
        "modules_mentioned": [
            "Motor Speed Control",
            "Zipper Merge",
            "Servo Gapper",
            "Print Apply",
            "Analog control module",
        ],
        "relationships_mentioned": [
            "SpdControl converts OffPercent/OnPercent to motor IO bit pattern",
            "Zipper Merge / Servo Gapper / Print Apply write OffPercent when Maint_IO is INVALID",
            "IO_Owner should match Conveyor process ownership when bits live in Conveyor",
        ],
        "reviewed": True,
        "notes": "Documents SpdControl table (menu Standard Controls | Motor Speed Control). Not present as SpdControl.asc in sample Greensboro RUN catalog.",
    },
    "FPC-Photoeye-Status-Reporting": {
        "relevance": "HIGH",
        "document_type": "module_ref",
        "generation_relevance": "IO",
        "subsystem": "photoeye",
        "tables_mentioned": [
            "PeList",
            "PeDisplay",
            "Batches",
            "Conveyor",
        ],
        "modules_mentioned": [
            "Photoeye Status Reporting",
            "PeReport Track(CTRL)",
            "PeReport Copy(GUI)",
            "PeReport Fill(ANY)",
        ],
        "relationships_mentioned": [
            "PeList selects monitored photoeyes; PeDisplay shows currently ON eyes ordered by duration",
            "Batches runs PeReport Track(CTRL) and PeReport Copy(GUI)",
            "Suggests Jamcheck_Sensor_Name as seed list for critical eyes",
        ],
        "reviewed": True,
        "notes": "Maintenance health reporting for photoeyes; useful for IO validation, not primary PLC codegen.",
    },
    "FPC-Ctrl-F4-Table-Distribution": {
        "relevance": "HIGH",
        "document_type": "procedure",
        "generation_relevance": "SITE_MODEL",
        "subsystem": "table_distribution",
        "tables_mentioned": [
            "Conveyor",
            "Mtrchain",
            "StartStopZones",
            "Jamzones",
            "Jamcheck",
            "Fullline",
            "Fulljam",
            "EStop",
            "Errors",
            "Machine",
            "Convclr",
            "PhView",
        ],
        "modules_mentioned": [
            "Ctrl-F4 Table Distribution",
            "Distributed Menus",
        ],
        "relationships_mentioned": [
            "Distributes selected configuration tables across ACs/machines",
            "Touches core site tables (Conveyor, zones, chains, errors) during multi-AC rollout",
        ],
        "reviewed": True,
        "notes": "Procedure for distributing FPC tables between machines; important for multi-AC site model consistency.",
    },
}


def load_asc_catalog(asc_dir: Path) -> dict[str, str]:
    """Map lower-case stem -> canonical stem from RUN ASC directory."""
    catalog: dict[str, str] = {}
    if not asc_dir.is_dir():
        return catalog
    for path in asc_dir.glob("*.asc"):
        stem = path.stem
        if stem.lower().startswith("old."):
            continue
        catalog[stem.lower()] = stem
    return catalog


def docx_text(path: Path, max_paras: int = 1500) -> str:
    """Extract paragraph + table-cell text from a .docx."""
    try:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        parts: list[str] = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                parts.append(p.text)
            if len(parts) >= max_paras:
                break
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        parts.append(t)
            if len(parts) >= max_paras + 800:
                break
        return "\n".join(parts)
    except Exception:
        return _docx_text_xml(path, max_paras=max_paras)


def _docx_text_xml(path: Path, max_paras: int = 1500) -> str:
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml")
    except Exception:
        return ""
    root = ET.fromstring(xml)
    paras: list[str] = []
    for p in root.iter(f"{W_NS}p"):
        texts = [t.text for t in p.iter(f"{W_NS}t") if t.text]
        if texts:
            paras.append("".join(texts))
        if len(paras) >= max_paras:
            break
    for tc in root.iter(f"{W_NS}tc"):
        texts = [t.text for t in tc.iter(f"{W_NS}t") if t.text]
        if texts:
            paras.append("".join(texts))
        if len(paras) >= max_paras + 800:
            break
    return "\n".join(paras)


def doc_strings(path: Path, limit: int = 200_000) -> str:
    """Best-effort ASCII strings from legacy .doc / other binaries."""
    try:
        data = path.read_bytes()
    except Exception:
        return ""
    strings = re.findall(rb"[\x20-\x7e]{5,}", data)
    text = "\n".join(s.decode("ascii", errors="ignore") for s in strings)
    return text[:limit]


def _looks_like_asc_stem(name: str) -> bool:
    """Heuristic: CamelCase, contains digit, or known underscore style."""
    if not name or len(name) < 4:
        return False
    key = name.lower()
    if key in REJECT_TABLE_NAMES:
        return False
    if "_" in name:
        return True
    if any(ch.isdigit() for ch in name):
        return True
    # CamelCase / PascalCase (e.g. SawMerge, HSSawLane, Fullline)
    if name[0].isupper() and any(ch.islower() for ch in name[1:]):
        return True
    if name.isupper() and len(name) >= 4:  # e.g. FORTNADT
        return True
    return False


def extract_tables(text: str, catalog: dict[str, str], filename_hints: list[str]) -> list[str]:
    """Extract ASC-like table names; prefer catalog hits; never invent."""
    found: dict[str, str] = {}
    hint_keys = {h.lower() for h in filename_hints}

    def add(name: str, require_catalog: bool = True) -> None:
        if not name or len(name) < 3:
            return
        key = name.lower()
        if key in REJECT_TABLE_NAMES and key not in hint_keys:
            return
        if key in catalog:
            if key in WEAK_ASC and key not in hint_keys:
                return
            found[catalog[key].lower()] = catalog[key]
            return
        if require_catalog:
            return
        if _looks_like_asc_stem(name) and key not in found:
            found[key] = name

    for hint in filename_hints:
        add(hint, require_catalog=False)

    for m in ASC_PAT.findall(text or ""):
        add(m, require_catalog=False)

    for m in TABLE_PAT.findall(text or ""):
        if m.lower() in catalog:
            add(m, require_catalog=True)
        elif _looks_like_asc_stem(m):
            add(m, require_catalog=False)

    # Catalog tokens on lines that mention table/.asc
    for line in (text or "").splitlines():
        low = line.lower()
        if "table" not in low and ".asc" not in low:
            continue
        for w in WORD_PAT.findall(line):
            if w.lower() in catalog and w.lower() not in WEAK_ASC:
                add(w, require_catalog=True)

    return sorted(found.values(), key=str.lower)


def infer_subsystem(stem: str, rel: str) -> str:
    s = stem.lower()
    if "pna" in s or "print" in s or "inkjet" in s or "zebra" in s or "p&a" in rel.lower():
        return "print_apply"
    if "sawtooth" in s or "hssaw" in s:
        return "merge/sawtooth"
    if "merge" in s or "zipper" in s:
        return "merge"
    if "sorter" in s or "shifter" in s or "pusher" in s:
        return "sorter"
    if "scanner" in s or "scan" in s:
        return "scanner"
    if "gap" in s:
        return "gapping"
    if "motor" in s or "mtrchain" in s or "spd" in s:
        return "motors"
    if "startstop" in s or "jam" in s or "full" in s or "estop" in s:
        return "zones"
    if (
        "iocard" in s
        or "fastio" in s
        or "fortnadt" in s
        or "configio" in s
        or "viewio" in s
        or "disable-bit" in s
        or s.endswith("-io")
        or "io-" in s
    ):
        return "io"
    if "msgmap" in s or "machine" in s or "fpcmsg" in s or "pc to pc" in s:
        return "messaging"
    if "transfer" in s or "xfer" in s:
        return "transfer"
    if "photoeye" in s or "pe" == s:
        return "photoeye"
    if "security" in s or "translation" in s or "menu" in s or "onscreen" in s:
        return "ui_admin"
    if "energy" in s or "ups" in s or "power" in s:
        return "power_energy"
    if "beacon" in s or "horn" in s or "timer" in s or "counter" in s or "trigger" in s:
        return "aux_controls"
    if "mysql" in s or "data-export" in s or "rest" in s or "opennms" in s:
        return "integrations"
    return "general"


def infer_document_type(stem: str) -> str:
    s = stem.lower()
    if "ddd" in s or "detail design" in s or "frd_" in s or "_idd" in s:
        return "ddd"
    if "checklist" in s:
        return "checklist"
    if "simulation" in s or "sim " in s or s.endswith("sim"):
        return "simulation"
    if "example" in s or "test" in s or "testplan" in s:
        return "example"
    if "procedure" in s or "upgrade" in s or "install" in s or "distribution" in s:
        return "procedure"
    if s.startswith("fpc-") or "module" in s or "configuration" in s:
        if any(k in s for k in ("how", "using", "setting", "setup", "viewio")):
            return "howto"
        return "module_ref"
    if "configuration" in s or "config" in s:
        return "howto"
    return "other"


def infer_relevance_and_gen(stem: str, subsystem: str) -> tuple[str, str]:
    s = stem.lower()
    # Explicit outs / historical
    if s.startswith("~$"):
        return "OUT_OF_SCOPE", "NONE"
    if "rev24" in s and "rev25" not in s:
        return "LOW", "HISTORICAL"
    if "old-transfers" in s or "old_transfers" in s:
        return "LOW", "HISTORICAL"
    if "translation" in s or "opennms" in s or "advanced-search" in s:
        return "LOW", "UI_ONLY"
    if "mysql" in s or "rest-api" in s or "data-export" in s:
        return "LOW", "NONE"
    if "security" in s or "menupad" in s or "onscreen" in s:
        return "LOW", "UI_ONLY"
    if subsystem == "print_apply":
        if "configuration" in s or "checklist" in s or "transfer" in s:
            return "MEDIUM", "PLC_GEN"
        return "LOW", "NONE"
    if any(
        k in s
        for k in (
            "merge-modules",
            "highspeedsawtooth",
            "sorter-control",
            "shifter-and-sorter",
            "startstopzones",
            "fulls-jams",
            "motor-startup",
        )
    ):
        return "CRITICAL", "PLC_GEN"
    if any(
        k in s
        for k in (
            "sorterconfigurationchecklist",
            "gapping-modules",
            "motor-speed",
            "scanner-control",
            "iocard",
            "fastio",
            "viewio",
            "fortnadt",
            "msgmap",
            "photoeye-status",
            "ctrl-f4",
            "unit sorter",
            "combinedjam",
            "initerrors",
            "disable-bits",
            "new-transfers",
        )
    ):
        if "ctrl-f4" in s or "msgmap" in s:
            gen = "SITE_MODEL"
        elif any(x in s for x in ("iocard", "fastio", "fortnadt", "viewio", "photoeye")):
            gen = "IO"
        elif "disable-bits" in s:
            gen = "IO"
        elif "initerrors" in s:
            gen = "PLC_GEN"
        else:
            gen = "PLC_GEN"
        return "HIGH", gen
    if subsystem in {"merge", "merge/sawtooth", "sorter", "zones", "motors", "gapping", "scanner", "transfer"}:
        return "MEDIUM", "PLC_GEN"
    if subsystem == "io":
        return "MEDIUM", "IO"
    if subsystem in {"messaging", "photoeye"}:
        return "MEDIUM", "SITE_MODEL"
    if subsystem == "aux_controls":
        return "MEDIUM", "PLC_GEN"
    # Generic FPC-* module/config docs are usually worth keeping at MEDIUM
    if s.startswith("fpc-") or s.startswith("fpcmsg"):
        if any(k in s for k in ("security", "translation", "menupad", "onscreen", "install", "mysql", "rest", "opennms", "data-export", "advanced-search")):
            return "LOW", "UI_ONLY" if any(k in s for k in ("security", "translation", "menu", "onscreen", "search")) else "NONE"
        if subsystem in {"integrations", "ui_admin", "power_energy"}:
            return "LOW", "NONE" if subsystem == "integrations" else ("UI_ONLY" if subsystem == "ui_admin" else "NONE")
        return "MEDIUM", "PLC_GEN"
    if any(k in s for k in ("sendconvimage", "convphs2png", "carousel", "pick to light")):
        return "LOW", "UI_ONLY" if "png" in s or "image" in s else "NONE"
    if any(k in s for k in ("beacon", "counter", "chainstretch", "analog", "housekeeping", "errors", "missing-pin", "inchstore", "endofwave", "horn", "energy", "packlift", "ratepoints", "specialflags", "timers", "trigger")):
        return "MEDIUM", "PLC_GEN"
    return "LOW", "NONE"


def filename_table_hints(stem: str) -> list[str]:
    hints: list[str] = []
    mapping = {
        "Merge-Modules": ["MergeInputs", "SawMerge", "SawLane"],
        "HighSpeedSawtoothMerge": ["HSSawMerge", "HSSawLane", "HSSawParm"],
        "StartStopZones": ["StartStopZones"],
        "Fulls-Jams-Fulljams": ["Fullline", "Fulljam", "Jamcheck"],
        "Motor-Startup-Chains": ["Mtrchain"],
        "FORTNADT": ["FORTNADT"],
        "MsgMap": ["MsgMap", "Machine"],
        "IOCard": ["IOCard"],
        "FastIO": ["FastIO"],
        "Gapping": ["NewGapper", "ServoGapper"],
        "Photoeye-Status": ["PeList", "PeDisplay"],
        "CombinedJamZones": ["CombinedJamZones", "CombinedEnableBits"],
    }
    for key, vals in mapping.items():
        if key.lower() in stem.lower():
            hints.extend(vals)
    return hints


def title_from_stem(stem: str) -> str:
    if stem.startswith("~$"):
        return f"(Office lock) {stem}"
    t = stem.replace("_", " ").replace("-", " ")
    return t


def iter_documents(training_root: Path) -> list[Path]:
    docs: list[Path] = []
    for corpus in CORPUS_DIRS:
        folder = training_root / corpus
        if not folder.is_dir():
            continue
        for path in sorted(folder.iterdir()):
            if not path.is_file():
                continue
            if path.name.startswith("~$"):
                continue
            if path.suffix.lower() in DOC_EXTS:
                docs.append(path)
    return docs


def list_zips(training_root: Path) -> list[dict]:
    zips = []
    for path in sorted(training_root.glob("*.zip")):
        zips.append(
            {
                "file": str(path.relative_to(REPO_ROOT)).replace("\\", "/"),
                "bytes": path.stat().st_size,
            }
        )
    return zips


def build_entry(
    path: Path,
    catalog: dict[str, str],
    skim_text: bool,
) -> dict:
    rel = str(path.relative_to(REPO_ROOT)).replace("\\", "/")
    stem = path.stem
    corpus = path.parent.name
    reviewed_meta = REVIEWED.get(stem)

    subsystem = (
        reviewed_meta.get("subsystem")
        if reviewed_meta
        else infer_subsystem(stem, rel)
    )
    document_type = (
        reviewed_meta.get("document_type")
        if reviewed_meta
        else infer_document_type(stem)
    )
    if reviewed_meta:
        relevance = reviewed_meta["relevance"]
        generation_relevance = reviewed_meta["generation_relevance"]
    else:
        relevance, generation_relevance = infer_relevance_and_gen(stem, subsystem)

    tables: list[str] = []
    modules: list[str] = []
    relationships: list[str] = []
    notes = ""
    reviewed = False
    text = ""

    if reviewed_meta:
        tables = list(reviewed_meta.get("tables_mentioned") or [])
        modules = list(reviewed_meta.get("modules_mentioned") or [])
        relationships = list(reviewed_meta.get("relationships_mentioned") or [])
        notes = reviewed_meta.get("notes") or ""
        reviewed = bool(reviewed_meta.get("reviewed"))
    elif skim_text:
        if path.suffix.lower() == ".docx":
            text = docx_text(path)
        elif path.suffix.lower() == ".doc":
            text = doc_strings(path)
        elif path.suffix.lower() == ".txt":
            text = path.read_text(encoding="utf-8", errors="replace")[:50_000]
        tables = extract_tables(text, catalog, filename_table_hints(stem))
        # light module hints from filename
        if "module" in stem.lower():
            modules = [title_from_stem(stem)]
        if tables:
            notes = "Tables inferred from filename/light text skim; not deep-reviewed."
        else:
            notes = "No ASC table names confirmed from light skim/filename heuristics."

    title = title_from_stem(stem)
    # Prefer human title from docx head only (legacy .doc strings are too noisy)
    if text and path.suffix.lower() == ".docx":
        for line in text.splitlines()[:20]:
            line = line.strip()
            if line.lower().startswith("fortnaplus"):
                title = re.sub(
                    r"^FortnaPlus\s*Control\s*", "", line, flags=re.I
                ).strip() or title
                break

    return {
        "title": title,
        "file": rel,
        "corpus": corpus,
        "extension": path.suffix.lower(),
        "bytes": path.stat().st_size,
        "subsystem": subsystem,
        "relevance": relevance,
        "document_type": document_type,
        "generation_relevance": generation_relevance,
        "tables_mentioned": tables,
        "modules_mentioned": modules,
        "relationships_mentioned": relationships,
        "reviewed": reviewed,
        "notes": notes,
    }


def polish_title(entry: dict, stem: str) -> None:
    """Use cleaner titles for known docs."""
    nicer = {
        "FPC-Merge-Modules": "FPC Merge Modules",
        "FPC-HighSpeedSawtoothMerge": "FPC High-Speed Sawtooth Merge",
        "FPC-Sorter-Control-Module": "FPC Sorter Control Module",
        "FPC-Shifter-And-Sorter-Configuration": "FPC Shifter And Sorter Configuration",
        "FPC-SorterConfigurationChecklist": "FPC Sorter Configuration Checklist",
        "Unit Sorter Control Software V1 rev25 DDD": "Unit Sorter Control Software V1 rev25 DDD",
        "Unit Sorter Control Software V1 rev24 DDD": "Unit Sorter Control Software V1 rev24 DDD",
        "FPC-StartStopZones": "FPC Start/Stop Zones",
        "FPC-Fulls-Jams-Fulljams": "FPC Fulls, Jams and Fulljams",
        "FPC-Motor-Startup-Chains": "FPC Motor Startup Chains",
        "FPC-The-ViewIO-Screen-and-FORTNADT-Table": "FPC View I/O Screen and FORTNADT Table",
        "FPC-Machine-MsgMap-Configuraton": "FPC Machine MsgMap Configuration (Configuraton spelling)",
        "FPC-IOCard-Interfaces": "FPC IOCard Interfaces",
        "FPC-FastIO-Configuration": "FPC FastIO Configuration",
        "FPC-Scanner-Control-Configuration": "FPC Scanner Control Configuration",
        "FPC-Gapping-Modules": "FPC Gapping Modules",
        "FPC-Motor-Speed-Control": "FPC Motor Speed Control",
        "FPC-Photoeye-Status-Reporting": "FPC Photoeye Status Reporting",
        "FPC-Ctrl-F4-Table-Distribution": "FPC Ctrl-F4 Table Distribution",
    }
    if stem in nicer:
        entry["title"] = nicer[stem]


def build_inventory(
    training_root: Path,
    asc_dir: Path,
    skim_all_docx: bool = True,
) -> dict:
    catalog = load_asc_catalog(asc_dir)
    docs = iter_documents(training_root)
    entries: list[dict] = []
    for path in docs:
        # Always skim non-reviewed docx lightly when enabled; reviewed use curated tables
        need_skim = skim_all_docx and path.stem not in REVIEWED
        entry = build_entry(path, catalog, skim_text=need_skim)
        polish_title(entry, path.stem)
        entries.append(entry)

    entries.sort(
        key=lambda e: (
            RELEVANCE_ORDER.get(e["relevance"], 99),
            e["subsystem"],
            e["file"],
        )
    )

    counts = Counter(e["relevance"] for e in entries)
    type_counts = Counter(e["document_type"] for e in entries)
    corpus_counts = Counter(e["corpus"] for e in entries)
    reviewed_count = sum(1 for e in entries if e["reviewed"])

    priority_stems = list(REVIEWED.keys())
    present = {Path(e["file"]).stem for e in entries}
    missing_priority = [s for s in priority_stems if s not in present]

    return {
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "training_root": str(training_root.relative_to(REPO_ROOT)).replace("\\", "/"),
        "asc_catalog_dir": (
            str(asc_dir.relative_to(REPO_ROOT)).replace("\\", "/")
            if asc_dir.exists()
            else str(asc_dir)
        ),
        "asc_catalog_size": len(catalog),
        "zips": list_zips(training_root),
        "counts": {
            "documents_total": len(entries),
            "by_relevance": dict(counts),
            "by_document_type": dict(type_counts),
            "by_corpus": dict(corpus_counts),
            "reviewed": reviewed_count,
            "critical": counts.get("CRITICAL", 0),
            "high": counts.get("HIGH", 0),
        },
        "missing_priority_docs": missing_priority,
        "documents": entries,
    }


def render_markdown(inv: dict) -> str:
    lines: list[str] = []
    lines.append("# FPC Training Document Index")
    lines.append("")
    lines.append(f"Generated: `{inv['generated_at']}`")
    lines.append("")
    lines.append("## Corpus")
    lines.append("")
    lines.append("Roots under `docs/training/`:")
    for c in CORPUS_DIRS:
        n = inv["counts"]["by_corpus"].get(c, 0)
        lines.append(f"- `{c}/` — {n} documents")
    lines.append("")
    lines.append("Zip archives present:")
    for z in inv["zips"]:
        lines.append(f"- `{z['file']}` ({z['bytes']} bytes)")
    pa_extra = REPO_ROOT / "docs" / "training" / "P&A Documents" / "PASIM1-RUN.tar.gz"
    if pa_extra.is_file():
        lines.append("")
        lines.append(
            f"Also present (not counted as a training document): "
            f"`docs/training/P&A Documents/PASIM1-RUN.tar.gz` ({pa_extra.stat().st_size} bytes)."
        )
    lines.append("")
    lines.append("## Summary")
    lines.append("")
    lines.append(f"- **Total documents:** {inv['counts']['documents_total']}")
    lines.append(f"- **CRITICAL:** {inv['counts']['critical']}")
    lines.append(f"- **HIGH:** {inv['counts']['high']}")
    lines.append(f"- **Reviewed (deep skim):** {inv['counts']['reviewed']}")
    if inv["missing_priority_docs"]:
        lines.append(
            "- **Missing priority docs:** "
            + ", ".join(f"`{s}`" for s in inv["missing_priority_docs"])
        )
    else:
        lines.append("- **Missing priority docs:** none")
    lines.append("")
    lines.append("Relevance breakdown:")
    for rel, n in sorted(
        inv["counts"]["by_relevance"].items(),
        key=lambda kv: RELEVANCE_ORDER.get(kv[0], 99),
    ):
        lines.append(f"- {rel}: {n}")
    lines.append("")
    lines.append("## CRITICAL / HIGH")
    lines.append("")
    lines.append(
        "Priority docs marked `reviewed=yes` were deep-skimmed; other HIGH rows "
        "come from filename/heuristics plus light docx skim."
    )
    lines.append("")
    lines.append(
        "| Relevance | Reviewed | Title | File | Subsystem | Gen | Tables (sample) |"
    )
    lines.append("|---|---|---|---|---|---|---|")
    for e in inv["documents"]:
        if e["relevance"] not in {"CRITICAL", "HIGH"}:
            continue
        tables = ", ".join(e["tables_mentioned"][:8])
        if len(e["tables_mentioned"]) > 8:
            tables += ", …"
        lines.append(
            f"| {e['relevance']} | {'yes' if e['reviewed'] else 'no'} | {e['title']} | "
            f"`{e['file']}` | {e['subsystem']} | {e['generation_relevance']} | {tables} |"
        )
    lines.append("")
    lines.append("## Full inventory")
    lines.append("")
    lines.append(
        "| Rel | Type | Gen | Subsystem | Title | File | Reviewed | Tables |"
    )
    lines.append("|---|---|---|---|---|---|---|---|")
    for e in inv["documents"]:
        tables = ", ".join(e["tables_mentioned"][:6])
        if len(e["tables_mentioned"]) > 6:
            tables += ", …"
        lines.append(
            f"| {e['relevance']} | {e['document_type']} | {e['generation_relevance']} | "
            f"{e['subsystem']} | {e['title']} | `{e['file']}` | "
            f"{'yes' if e['reviewed'] else 'no'} | {tables} |"
        )
    lines.append("")
    lines.append("## Notes")
    lines.append("")
    lines.append(
        "- `tables_mentioned` prefers exact ASC-like names from documents; "
        "validated against `workspace/active/RUN/FORTNA` when present."
    )
    lines.append(
        "- Office lock files (`~$*`) are excluded."
    )
    lines.append(
        "- Rebuild: `python tools/scripts/fortna_fpc_document_inventory.py`"
    )
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--training-root", type=Path, default=TRAINING_ROOT)
    ap.add_argument("--asc-dir", type=Path, default=DEFAULT_ASC_DIR)
    ap.add_argument("--json-out", type=Path, default=DEFAULT_JSON)
    ap.add_argument("--md-out", type=Path, default=DEFAULT_MD)
    ap.add_argument(
        "--no-skim",
        action="store_true",
        help="Skip light text skim of non-reviewed docx (filename heuristics only)",
    )
    args = ap.parse_args()

    inv = build_inventory(
        training_root=args.training_root,
        asc_dir=args.asc_dir,
        skim_all_docx=not args.no_skim,
    )

    args.json_out.parent.mkdir(parents=True, exist_ok=True)
    args.md_out.parent.mkdir(parents=True, exist_ok=True)
    args.json_out.write_text(json.dumps(inv, indent=2) + "\n", encoding="utf-8")
    args.md_out.write_text(render_markdown(inv), encoding="utf-8")

    print(f"Wrote {args.json_out}")
    print(f"Wrote {args.md_out}")
    print(
        f"documents={inv['counts']['documents_total']} "
        f"CRITICAL={inv['counts']['critical']} HIGH={inv['counts']['high']} "
        f"reviewed={inv['counts']['reviewed']}"
    )
    if inv["missing_priority_docs"]:
        print("missing_priority:", ", ".join(inv["missing_priority_docs"]))
    else:
        print("missing_priority: none")


if __name__ == "__main__":
    main()
